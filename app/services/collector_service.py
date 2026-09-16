from __future__ import annotations

import hashlib
import json
import re
import shutil
from pathlib import Path
from urllib.parse import urljoin, urlsplit, urlunsplit, parse_qsl, urlencode
from typing import Any

from app.database.repositories.core import SourceRepository, DocumentRepository, ImageRepository


class CollectorService:
    """本地资料采集服务。3.2 增强文本元数据识别、统计、结果详情与异常分类。"""

    TEXT_EXTENSIONS = {".txt", ".md", ".csv", ".json", ".log"}
    IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".gif", ".tif", ".tiff"}

    def __init__(self, db_path, data_root: Path):
        self.db_path = Path(db_path)
        self.data_root = Path(data_root)
        self.import_root = self.data_root / "imports"
        self.image_root = self.data_root / "knowledge" / "images"
        self.import_root.mkdir(parents=True, exist_ok=True)
        self.image_root.mkdir(parents=True, exist_ok=True)
        self.sources = SourceRepository(self.db_path)
        self.documents = DocumentRepository(self.db_path)
        self.images = ImageRepository(self.db_path)

    @staticmethod
    def hash_bytes(data: bytes) -> str:
        return hashlib.sha256(data).hexdigest()

    @staticmethod
    def hash_file(path: Path, chunk_size: int = 1024 * 1024) -> str:
        h = hashlib.sha256()
        with path.open("rb") as f:
            while chunk := f.read(chunk_size):
                h.update(chunk)
        return h.hexdigest()

    @staticmethod
    def detect_language(text: str) -> str:
        """轻量语言识别：不引入模型依赖，供后续 NLP/LLM 阶段替换。"""
        if not text.strip():
            return "unknown"
        zh = len(re.findall(r"[\u4e00-\u9fff]", text))
        en = len(re.findall(r"[A-Za-z]", text))
        if zh == 0 and en == 0:
            return "other"
        if zh >= en * 0.15 and zh > 0:
            return "zh" if zh >= en else "zh-en"
        return "en"

    @staticmethod
    def word_count(text: str) -> int:
        # 中文按字符计数，英文/数字按词计数；过滤标点和空白。
        chinese = re.findall(r"[\u4e00-\u9fff]", text)
        non_chinese = re.findall(r"[A-Za-z0-9]+(?:[-_][A-Za-z0-9]+)*", text)
        return len(chinese) + len(non_chinese)

    @staticmethod
    def char_count(text: str) -> int:
        return len(re.sub(r"\s", "", text))

    @staticmethod
    def infer_title(text: str, fallback: str = "") -> str:
        """从 Markdown 标题/首个有效行推断标题，失败时使用文件名或正文前缀。"""
        for line in text.splitlines():
            line = line.strip().lstrip("#").strip()
            if line and len(line) <= 120:
                return line
        return fallback.strip() or text.strip()[:80]

    @staticmethod
    def infer_content_type(text: str) -> str:
        stripped = text.strip()
        if re.search(r"(?:^|\n)\s*(?:positive prompt|negative prompt|prompt)\s*[:：]", stripped, re.I):
            return "prompt"
        if stripped.startswith("{") or stripped.startswith("["):
            try:
                json.loads(stripped)
                return "json"
            except Exception:
                pass
        if stripped.startswith("#"):
            return "markdown"
        return "text"

    @classmethod
    def build_text_metadata(cls, text: str, title: str = "") -> dict[str, Any]:
        clean = (text or "").replace("\r\n", "\n").replace("\r", "\n").strip()
        inferred = cls.infer_title(clean, title)
        return {
            "title": inferred,
            "language": cls.detect_language(clean),
            "word_count": cls.word_count(clean),
            "char_count": cls.char_count(clean),
            "content_type": cls.infer_content_type(clean),
            "line_count": len(clean.splitlines()),
        }

    def _find_source_by_hash(self, content_hash: str):
        return self.sources.list(1, 0, "content_hash=?", (content_hash,))

    def _find_document_by_hash(self, content_hash: str):
        return self.documents.list(1, 0, "content_hash=?", (content_hash,))

    def _find_image_by_hash(self, file_hash: str):
        return self.images.list(1, 0, "file_hash=?", (file_hash,))

    def collect_text(self, text: str, title: str = "", source_type: str = "manual_text", description: str = "") -> dict[str, Any]:
        text = (text or "").replace("\r\n", "\n").replace("\r", "\n").strip()
        if not text:
            raise ValueError("采集文本不能为空")
        meta = self.build_text_metadata(text, title)
        content_hash = self.hash_bytes(text.encode("utf-8"))
        existing = self._find_source_by_hash(content_hash) or self._find_document_by_hash(content_hash)
        if existing:
            row = existing[0]
            return {
                "status": "duplicate", "source_id": row.get("source_id", row.get("id")),
                "document_id": row.get("id") if "content" in row else None,
                "message": "内容已存在，已跳过重复入库", "metadata": meta,
            }
        source_id = self.sources.create({
            "title": meta["title"], "source_type": source_type,
            "description": description or f"{meta['content_type']}文本采集 · {meta['language']}",
            "content_hash": content_hash, "status": "completed",
        })
        document_id = self.documents.create({
            "source_id": source_id, "title": meta["title"], "content": text,
            "content_hash": content_hash, "language": meta["language"], "word_count": meta["word_count"],
            "summary": text[:300],
        })
        return {
            "status": "created", "source_id": source_id, "document_id": document_id,
            "message": "文本采集成功", "metadata": meta,
        }

    def collect_file(self, file_path: str | Path, title: str = "", copy_to_library: bool = True) -> dict[str, Any]:
        path = Path(file_path).expanduser().resolve()
        if not path.exists() or not path.is_file():
            raise FileNotFoundError(f"文件不存在：{path}")
        suffix = path.suffix.lower()
        try:
            if suffix in self.IMAGE_EXTENSIONS:
                return self._collect_image(path, title, copy_to_library)
            if suffix in self.TEXT_EXTENSIONS:
                return self._collect_text_file(path, title)
            if suffix == ".docx":
                return self._collect_docx(path, title)
            if suffix == ".pdf":
                return self._collect_pdf(path, title)
        except Exception as exc:
            return {"status": "failed", "message": str(exc), "file_path": str(path), "file_name": path.name}
        return {"status": "failed", "message": f"暂不支持的文件类型：{suffix or '无扩展名'}", "file_path": str(path), "file_name": path.name}

    def _collect_text_file(self, path: Path, title: str) -> dict[str, Any]:
        raw = path.read_bytes()
        digest = self.hash_bytes(raw)
        existing = self._find_source_by_hash(digest)
        if existing:
            return {"status": "duplicate", "source_id": existing[0]["id"], "document_id": None, "message": "文件内容已存在，已跳过重复导入", "file_name": path.name}
        text = None
        encoding_used = "utf-8"
        for encoding in ("utf-8-sig", "utf-8", "gb18030", "big5"):
            try:
                text = raw.decode(encoding)
                encoding_used = encoding
                break
            except UnicodeDecodeError:
                pass
        if text is None:
            text = raw.decode("utf-8", errors="replace")
            encoding_used = "utf-8(replace)"
        result = self.collect_text(text, title or path.stem, "file_text", f"本地文件：{path.name}；编码：{encoding_used}")
        result.update({"file_name": path.name, "file_path": str(path), "encoding": encoding_used})
        return result

    def _collect_docx(self, path: Path, title: str) -> dict[str, Any]:
        from docx import Document
        doc = Document(path)
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        text = "\n".join(paragraphs)
        table_count = len(doc.tables)
        for table in doc.tables:
            rows = []
            for row in table.rows:
                rows.append(" | ".join(cell.text.strip() for cell in row.cells))
            if rows:
                text += "\n" + "\n".join(rows)
        if not text.strip():
            raise ValueError("DOCX中没有可采集的文字内容")
        result = self.collect_text(text, title or path.stem, "docx", f"Word文件：{path.name}；表格：{table_count}个")
        result.update({"file_name": path.name, "file_path": str(path), "table_count": table_count})
        return result

    def _collect_pdf(self, path: Path, title: str) -> dict[str, Any]:
        import fitz
        pages = []
        page_count = 0
        with fitz.open(path) as pdf:
            page_count = len(pdf)
            for page in pdf:
                pages.append(page.get_text("text"))
        text = "\n".join(pages).strip()
        if not text:
            raise ValueError("PDF没有可直接提取的文字内容；扫描PDF将在后续OCR阶段处理")
        result = self.collect_text(text, title or path.stem, "pdf", f"PDF文件：{path.name}；页数：{page_count}")
        result.update({"file_name": path.name, "file_path": str(path), "page_count": page_count})
        return result

    def _collect_image(self, path: Path, title: str, copy_to_library: bool) -> dict[str, Any]:
        file_hash = self.hash_file(path)
        existing = self._find_image_by_hash(file_hash)
        if existing:
            return {"status": "duplicate", "source_id": existing[0]["source_id"], "image_id": existing[0]["id"], "message": "图片已存在，已跳过重复导入", "file_name": path.name}
        target = path
        if copy_to_library:
            target = self.image_root / f"{file_hash[:16]}{path.suffix.lower()}"
            if not target.exists():
                shutil.copy2(path, target)
        width = height = None
        image_format = path.suffix.lower().lstrip(".")
        try:
            from PIL import Image
            with Image.open(path) as im:
                width, height = im.size
                image_format = (im.format or image_format).lower()
        except Exception:
            pass
        source_id = self.sources.create({
            "title": title.strip() or path.stem, "source_type": "image",
            "description": f"图片文件：{path.name}", "content_hash": file_hash, "status": "completed",
        })
        image_id = self.images.create({
            "source_id": source_id, "file_path": str(target), "original_url": None,
            "file_hash": file_hash, "width": width, "height": height, "format": image_format,
            "metadata": json.dumps({"original_path": str(path)}, ensure_ascii=False),
            "analysis_status": "pending",
        })
        return {
            "status": "created", "source_id": source_id, "image_id": image_id,
            "message": "图片采集成功", "file_path": str(target), "file_name": path.name,
            "metadata": {"width": width, "height": height, "format": image_format, "file_hash": file_hash},
        }

    def find_source_by_url(self, url: str):
        """按 URL 查找已采集来源（供“查看采集结果同步当前网址”使用）。"""
        try:
            normalized = self.normalize_url(url)
        except ValueError:
            return None
        rows = self._find_source_by_url(normalized)
        return rows[0] if rows else None

    def recent_sources(self, limit=30):
        """最近的采集来源，供采集历史列表使用。"""
        return self.sources.list(limit=limit)

    def load_source_content(self, source_id):
        """加载来源对应的文档内容或图片信息，供采集历史预览。"""
        source = self.sources.get(source_id)
        if not source:
            return None
        doc_rows = self.documents.list(1, 0, "source_id=?", (source_id,))
        if doc_rows:
            doc = doc_rows[0]
            return {"kind": "document", "source": source, "title": doc.get("title") or source.get("title"),
                    "content": doc.get("content") or "", "language": doc.get("language"),
                    "word_count": doc.get("word_count")}
        img_rows = self.images.list(1, 0, "source_id=?", (source_id,))
        if img_rows:
            img = img_rows[0]
            return {"kind": "image", "source": source, "title": source.get("title"),
                    "content": json.dumps({"file_path": img.get("file_path"), "width": img.get("width"),
                                           "height": img.get("height"), "format": img.get("format"),
                                           "metadata": img.get("metadata")}, ensure_ascii=False, indent=2)}
        return {"kind": "empty", "source": source, "title": source.get("title"), "content": ""}

    def collect(self, kind: str, **kwargs) -> dict[str, Any]:
        if kind == "text":
            return self.collect_text(kwargs.get("text", ""), kwargs.get("title", ""))
        if kind == "file":
            return self.collect_file(kwargs["file_path"], kwargs.get("title", ""))
        if kind == "url":
            return self.collect_url(kwargs.get("url", ""), kwargs.get("title", ""), kwargs.get("timeout", 20.0))
        raise ValueError(f"未知采集类型：{kind}")

    @staticmethod
    def parse_text_structure(text: str) -> dict[str, Any]:
        lines = (text or '').replace('\r\n','\n').replace('\r','\n').splitlines()
        headings = []
        paragraphs = []
        current = []
        for idx, raw in enumerate(lines, 1):
            line = raw.strip()
            if not line:
                if current:
                    paragraphs.append({'text': '\n'.join(current), 'start_line': idx-len(current), 'end_line': idx-1})
                    current=[]
                continue
            m = re.match(r'^(#{1,6})\s+(.+?)\s*#*$', line)
            if m:
                if current:
                    paragraphs.append({'text': '\n'.join(current), 'start_line': idx-len(current), 'end_line': idx-1}); current=[]
                headings.append({'level': len(m.group(1)), 'text': m.group(2).strip(), 'line': idx})
            else:
                current.append(line)
        if current:
            paragraphs.append({'text': '\n'.join(current), 'start_line': len(lines)-len(current)+1, 'end_line': len(lines)})
        return {'line_count': len(lines), 'heading_count': len(headings), 'paragraph_count': len(paragraphs), 'headings': headings, 'paragraphs': paragraphs}

    @staticmethod
    def parse_docx_structure(path: str | Path) -> dict[str, Any]:
        from docx import Document
        doc = Document(path)
        blocks = []
        for p in doc.paragraphs:
            text = p.text.strip()
            if text:
                style = p.style.name if p.style else ''
                level = 0
                m = re.search(r'(?:Heading|标题)\s*(\d+)', style, re.I)
                if m: level = int(m.group(1))
                blocks.append({'type': 'heading' if level else 'paragraph', 'level': level, 'text': text, 'style': style})
        tables=[]
        for table in doc.tables:
            tables.append([[cell.text.strip() for cell in row.cells] for row in table.rows])
        return {'paragraph_count': len(doc.paragraphs), 'nonempty_block_count': len(blocks), 'table_count': len(tables), 'blocks': blocks, 'tables': tables}

    @staticmethod
    def parse_pdf_structure(path: str | Path) -> dict[str, Any]:
        import fitz
        pages=[]
        with fitz.open(path) as pdf:
            for number, page in enumerate(pdf, 1):
                text=page.get_text('text') or ''
                pages.append({'page': number, 'text': text.strip(), 'char_count': len(text.strip()), 'block_count': len(page.get_text('blocks'))})
        return {'page_count': len(pages), 'nonempty_page_count': sum(1 for p in pages if p['text']), 'pages': pages}

    @staticmethod
    def normalize_url(url: str) -> str:
        value = (url or '').strip()
        if not value:
            raise ValueError('URL不能为空')
        if not re.match(r'^https?://', value, re.I):
            value = 'https://' + value
        parts = urlsplit(value)
        if not parts.netloc:
            raise ValueError('URL格式不正确')
        # 去除常见追踪参数，使同一网页的分享链接能够命中去重。
        tracking = {'utm_source', 'utm_medium', 'utm_campaign', 'utm_term', 'utm_content', 'spm', 'from'}
        query = [(k, v) for k, v in parse_qsl(parts.query, keep_blank_values=True) if k.lower() not in tracking]
        return urlunsplit((parts.scheme.lower(), parts.netloc.lower(), parts.path or '/', urlencode(query), ''))

    def _find_source_by_url(self, url: str):
        return self.sources.list(1, 0, 'url=?', (url,))

    def _download_web_images(self, soup, base_url: str, client, max_images: int = 30) -> dict[str, Any]:
        results = []
        seen = set()
        image_dir = self.image_root / 'web'
        image_dir.mkdir(parents=True, exist_ok=True)
        for tag in soup.find_all('img'):
            src = (tag.get('src') or tag.get('data-src') or tag.get('data-original') or '').strip()
            if not src or src.startswith(('data:', 'javascript:')):
                continue
            image_url = urljoin(base_url, src)
            if image_url in seen or not image_url.lower().startswith(('http://', 'https://')):
                continue
            seen.add(image_url)
            if len(results) >= max_images:
                break
            try:
                response = client.get(image_url)
                response.raise_for_status()
                content = response.content
                if not content:
                    continue
                file_hash = self.hash_bytes(content)
                existing = self._find_image_by_hash(file_hash)
                if existing:
                    results.append({'status': 'duplicate', 'image_id': existing[0]['id'], 'url': image_url})
                    continue
                suffix = Path(image_url.split('?', 1)[0]).suffix.lower()
                if suffix not in self.IMAGE_EXTENSIONS:
                    suffix = '.jpg'
                target = image_dir / f'{file_hash[:24]}{suffix}'
                target.write_bytes(content)
                width = height = None
                image_format = suffix.lstrip('.')
                try:
                    from PIL import Image
                    with Image.open(target) as im:
                        width, height = im.size
                        image_format = (im.format or image_format).lower()
                except Exception:
                    pass
                results.append({'status': 'downloaded', 'url': image_url, 'file_path': str(target),
                                'file_hash': file_hash, 'width': width, 'height': height, 'format': image_format})
            except Exception as exc:
                results.append({'status': 'failed', 'url': image_url, 'message': str(exc)})
        return {'total': len(results), 'downloaded': sum(1 for x in results if x['status'] == 'downloaded'),
                'duplicate': sum(1 for x in results if x['status'] == 'duplicate'),
                'failed': sum(1 for x in results if x['status'] == 'failed'), 'items': results}

    def collect_url(self, url: str, title: str = '', timeout: float = 20.0, download_images: bool = True) -> dict[str, Any]:
        """采集普通网页及微信公众号文章等HTML页面。"""
        normalized = self.normalize_url(url)
        existing = self._find_source_by_url(normalized)
        if existing:
            return {'status': 'duplicate', 'source_id': existing[0]['id'], 'url': normalized,
                    'message': '该网页URL已经采集过，已跳过重复导入'}
        try:
            import httpx
            from bs4 import BeautifulSoup
            headers = {'User-Agent': 'Mozilla/5.0 (PromptForge Collector)'}
            with httpx.Client(follow_redirects=True, timeout=timeout, headers=headers) as client:
                response = client.get(normalized)
                response.raise_for_status()
                html = response.text
                final_url = str(response.url)
                soup = BeautifulSoup(html, 'html.parser')
                image_result = self._download_web_images(soup, final_url, client) if download_images else {'total': 0, 'downloaded': 0, 'duplicate': 0, 'failed': 0, 'items': []}
            # HTML 已在请求上下文中解析，保留 final_url 与 soup。
            for node in soup(['script', 'style', 'noscript', 'svg']):
                node.decompose()
            page_title = title.strip() or (soup.title.get_text(' ', strip=True) if soup.title else '')
            if not page_title:
                og = soup.find('meta', attrs={'property': 'og:title'})
                page_title = (og.get('content', '').strip() if og else '') or final_url
            author = ''
            for key in ('author', 'article:author'):
                tag = soup.find('meta', attrs={'name': key}) or soup.find('meta', attrs={'property': key})
                if tag and tag.get('content'):
                    author = tag['content'].strip(); break
            publish_time = ''
            for key in ('article:published_time', 'publish_date', 'date'):
                tag = soup.find('meta', attrs={'property': key}) or soup.find('meta', attrs={'name': key})
                if tag and tag.get('content'):
                    publish_time = tag['content'].strip(); break
            article = soup.find('article') or soup.find('main') or soup.body
            text = article.get_text('\n', strip=True) if article else ''
            text = re.sub(r'\n{3,}', '\n\n', text).strip()
            if not text:
                raise ValueError('网页未提取到正文内容')
            content_hash = self.hash_bytes(text.encode('utf-8'))
            duplicate_content = self._find_source_by_hash(content_hash) or self._find_document_by_hash(content_hash)
            if duplicate_content:
                return {'status': 'duplicate', 'url': final_url, 'source_id': duplicate_content[0].get('source_id', duplicate_content[0].get('id')), 'document_id': duplicate_content[0].get('id') if 'content' in duplicate_content[0] else None, 'message': '网页正文内容已存在，已跳过重复入库'}
            raw_dir = self.import_root / 'html'
            raw_dir.mkdir(parents=True, exist_ok=True)
            raw_path = raw_dir / f'{self.hash_bytes(final_url.encode())[:24]}.html'
            raw_path.write_text(html, encoding='utf-8')
            source_id = self.sources.create({'title': page_title, 'url': final_url, 'source_type': 'web',
                'author': author, 'publish_time': publish_time, 'description': '网页采集',
                'content_hash': content_hash, 'status': 'completed'})
            document_id = self.documents.create({'source_id': source_id, 'title': page_title, 'content': text,
                'content_hash': content_hash, 'language': self.detect_language(text),
                'word_count': self.word_count(text), 'summary': text[:300], 'raw_html_path': str(raw_path)})
            for item in image_result.get('items', []):
                if item.get('status') != 'downloaded':
                    continue
                self.images.create({'source_id': source_id, 'document_id': document_id,
                    'file_path': item['file_path'], 'original_url': item['url'],
                    'file_hash': item['file_hash'], 'width': item.get('width'),
                    'height': item.get('height'), 'format': item.get('format'),
                    'metadata': json.dumps({'source_url': final_url}, ensure_ascii=False),
                    'analysis_status': 'pending'})
            final_status = 'partial' if image_result.get('failed', 0) else 'created'
            final_message = '网页正文采集成功，但部分图片下载失败' if final_status == 'partial' else '网页采集成功'
            return {'status': final_status, 'source_id': source_id, 'document_id': document_id,
                'url': final_url, 'message': final_message,
                'metadata': {**self.build_text_metadata(text, page_title), 'author': author,
                             'publish_time': publish_time, 'raw_html_path': str(raw_path), 'images': image_result}}
        except Exception as exc:
            return {'status': 'failed', 'url': normalized, 'message': f'网页采集失败：{exc}'}
