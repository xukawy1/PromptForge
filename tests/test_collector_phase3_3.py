from app.services.collector_service import CollectorService

def test_text_structure():
    r=CollectorService.parse_text_structure('# 标题\n\n第一段\n第二行\n\n## 子标题\n内容')
    assert r['heading_count']==2
    assert r['paragraph_count']==2

def test_docx_structure(tmp_path):
    from docx import Document
    p=tmp_path/'a.docx'; d=Document(); d.add_heading('主标题',1); d.add_paragraph('正文'); t=d.add_table(rows=1, cols=2); t.cell(0,0).text='A'; t.cell(0,1).text='B'; d.save(p)
    r=CollectorService.parse_docx_structure(p)
    assert r['table_count']==1 and r['nonempty_block_count']==2

def test_pdf_structure(tmp_path):
    import fitz
    p=tmp_path/'a.pdf'; doc=fitz.open(); page=doc.new_page(); page.insert_text((72,72),'PDF正文'); doc.save(p); doc.close()
    r=CollectorService.parse_pdf_structure(p)
    assert r['page_count']==1 and r['nonempty_page_count']==1
